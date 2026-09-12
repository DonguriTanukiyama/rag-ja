# -*- coding: utf-8 -*-
u"""genba_norm.py — 現場の帳票を、RAG が読める Markdown に正規化する

    python genba_norm.py <入力ディレクトリ> -o <出力ディレクトリ>
    python genba_norm.py <入力> -o <出力> --single-file   点検記録を1ファイルにまとめる
    python genba_norm.py <入力> -o <出力> --naive         正規化しない対照（測定用）
    python genba_norm.py --demo                           自己確認

`rag_ja.py` は .md / .txt / .rst しか読まない。Excel の点検記録は食えない。
ここで .md に落として、そのまま `python rag_ja.py index <出力ディレクトリ>` に渡す。

------------------------------------------------------------------
なぜ「そのまま食わせる」ではいけないか
------------------------------------------------------------------

  ① 入力が汚い
     1つの日付列に6通りの書き方が混ざる。Excel が勝手にシリアル値へ変える。
     日付セルが縦に結合されていて、3行のうち2行は空になる。
     人名・部署名・設備名が揺れる。数値に単位と全角が混ざる。

  ② **「異常なし」と空欄の区別がつかない**
     点検欄が空なのは、異常が無かったからではなく、**点検していないから**のことがある。
     空欄を黙って落とすと、残るのは「異常なし」だけになり、
     「今月は異常が無かった」と読めてしまう。**いちばん危ない壊れ方。**

  ③ 答えられない質問に答えてしまう
     → これは rag_ja.py 側で既に対処済み（出典の無い文を出さない）。

このファイルが埋めるのは ① と ②。

------------------------------------------------------------------
② をどう扱うか — 点検欄を6つに分ける
------------------------------------------------------------------

  記入あり        異常なし / 特になし / OK / レ / ✓ / 良 / ○
  異常            異常あり… / NG / 不良
  **対象外**      計画停止 / 停止中 / 休転 / 非稼働 / 対象外
                  → **空欄でも「未記入」ではない。**設備が止まっていれば点検しない
  **条件付き許容** 要観察 / 経過観察 / 軽微 / 稼働可 / 要確認
                  → **異常ではない。**ここを異常に入れると異常件数が実態より膨らむ
  **未記入**      空欄。→ 本文に「未記入」と書く。**黙って消さない**
  **判別不能**    - / ー / ／ / N/A
                  「該当なし」なのか「未チェック」なのか、原票からは決められない

**決められないことを、決められないと書く。**

2026-09-12 のレビューで「対象外」「条件付き許容」が抜けていると指摘され、4分類から6分類にした。
それまで `要観察` を **異常** に数えていた。異常件数が実態より多く出る誤り。

そして **集計は前処理で作って文書に書く。**
RAG は検索であって集計ではない。「未記入は何件か」に答えさせたいなら、
その一文がどこかに書かれている必要がある。だからサマリ文書を1本出す。

------------------------------------------------------------------
xlsx の読み取り
------------------------------------------------------------------

**openpyxl があれば使う。無ければ内蔵の読み取り（標準ライブラリのみ）に落ちる。**
どちらを使ったかは必ず表示する。**黙って弱いほうに落ちない**（rag_ja.py の埋め込みと同じ方針）。

内蔵の読み取りは zip + XML を自前で解いている。次に未対応:
  - 1904年日付システム（Mac 由来）
  - Strict ネームスペース
  - 旧 .xls（OLE複合ドキュメント）
  - 複数シート（`sheet1.xml` 決め打ち）

`--demo` は、openpyxl がある環境では **内蔵の読み取りと openpyxl の結果を突き合わせる。**

------------------------------------------------------------------
front-matter
------------------------------------------------------------------

各文書の先頭に `date` / `equipment` / `person` / `dept` などを出す。
`rag_ja.py` 側がこれを metadata に入れ、`--where date=2026-09-03` で絞れる。
**「9月3日の」のような条件は、意味の近さではなく完全一致で絞るべきもの。**

依存: なし（openpyxl があれば使うが、無くても動く）
"""
import argparse
import datetime
import io
import os
import re
import sys
import unicodedata
import zipfile
import xml.etree.ElementTree as ET

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"

# 表記ゆれ。左に寄せる。**現場ごとに違うので、ここは必ず書き換えて使う。**
ALIASES = {
    u"設備": {
        u"加圧浮上槽": [u"浮上槽", u"加圧浮上", u"加圧浮上そう"],
        u"活性汚泥槽": [u"汚泥槽", u"活性汚泥"],
    },
    u"部署": {
        u"第1製造部": [u"第一製造部", u"製造1部", u"製造一部", u"M1", u"Ｍ１"],
    },
    u"担当": {
        u"山田太郎": [u"山田", u"山田 太郎", u"ヤマダ", u"やまだ"],
        u"佐藤花子": [u"佐藤", u"佐藤 花子", u"さとう", u"サトウ"],
    },
}

OK_MARKS = {u"異常なし", u"異常無し", u"特になし", u"特に無し",
            u"OK", u"ok", u"レ", u"✓", u"良", u"可", u"○", u"〇"}
AMBIG_MARKS = {u"-", u"ー", u"−", u"—", u"/", u"／", u"N/A", u"n/a", u"NA"}

# 接頭辞で見る。**判定の順番が本体。**下の classify() を読むこと
EXCLUDED_HEAD = (u"対象外", u"点検対象外", u"計画停止", u"停止中", u"休転", u"非稼働", u"停止")
COND_HEAD = (u"要観察", u"経過観察", u"継続監視", u"軽微", u"稼働可", u"要確認", u"要注意")
ABNORMAL_HEAD = (u"異常", u"NG", u"不良", u"要修理", u"故障")

RECORDED = u"記入あり"
ABNORMAL = u"異常"
EXCLUDED = u"対象外"
CONDITIONAL = u"条件付き許容"
BLANK = u"未記入"
AMBIGUOUS = u"判別不能"

KINDS = (RECORDED, ABNORMAL, CONDITIONAL, EXCLUDED, BLANK, AMBIGUOUS)

# 「点検されたという記録がある」とみなす状態。記入率の分子
COUNTED_AS_CHECKED = (RECORDED, ABNORMAL, CONDITIONAL, EXCLUDED)


# --- 文字をそろえる ---------------------------------------------------------

def z2h(s):
    u"""全角の英数記号を半角に。日本語はそのまま。"""
    return unicodedata.normalize("NFKC", s) if s else s


def canon_text(kind, s):
    u"""**本文中**の別名を代表名に寄せる。セル1つぶんの canon() とは別物。

    2026-09-12 に実データで出た誤り:
      「加圧浮上槽」に対して 別名「加圧浮上」→「加圧浮上槽」を当て、
      さらに 別名「浮上槽」→「加圧浮上槽」を当ててしまい、
      **「加圧加圧浮上槽槽」**になっていた。**既に正しい形にも置換が当たる。**
      --demo では出ない。本文を目で見て気づいた種類の誤り。
    → 代表名と別名をいったん退避してから、まとめて戻す。二度当たらない。
    """
    SEP = chr(0)
    for rep, alts in ALIASES.get(kind, {}).items():
        s = s.replace(rep, SEP)                              # 既に正しい形を先に退避
        for a in sorted(alts, key=len, reverse=True):
            s = s.replace(a, SEP)
        s = s.replace(SEP, rep)
    return s


def canon(kind, value):
    u"""表記ゆれを代表名に寄せる。表に無ければそのまま返す。"""
    v = z2h((value or u"").strip())
    table = ALIASES.get(kind, {})
    for rep, alts in table.items():
        if v == rep or v in alts or z2h(rep) == v:
            return rep
    return v


def find_equipment(text):
    u"""本文に出てくる設備名を代表名で拾う。メタデータ用。"""
    t = canon_text(u"設備", text)
    return [rep for rep in ALIASES.get(u"設備", {}) if rep in t]


# --- 日付 -------------------------------------------------------------------

def serial_to_date(n):
    u"""Excel の1900年日付システム。1900年うるう年バグぶんを含めて 1899-12-30 起点。

    **1904年日付システム（Mac 由来）には未対応。**4年ずれる。
    """
    return datetime.date(1899, 12, 30) + datetime.timedelta(days=int(n))


def parse_date(raw, default_year=None, default_month=None):
    u"""1つの列に混ざった書き方を全部受ける。読めなければ None を返す（捨てない）。

    受ける形: 2026/9/1  2026-09-01  令和8年9月3日  R8.9.5  26.9.6  9/4  シリアル値

    **既知の誤検知**: `12.3.4` のようなバージョン番号を 2012-03-04 と読む。
    日付列と日報の1行目からしか呼んでいないので実害は出ていないが、穴ではある。
    """
    if raw is None:
        return None
    s = z2h(str(raw)).strip()
    if not s:
        return None

    if re.match(r"^\d{5}(\.\d+)?$", s):
        try:
            return serial_to_date(float(s))
        except Exception:
            return None

    m = re.match(r"^(令和|R|r)\s*(\d{1,2})[.年/\-](\d{1,2})[.月/\-](\d{1,2})日?$", s)
    if m:
        return _mk(2018 + int(m.group(2)), int(m.group(3)), int(m.group(4)))

    m = re.match(r"^(\d{4})[/\-.年](\d{1,2})[/\-.月](\d{1,2})日?$", s)
    if m:
        return _mk(int(m.group(1)), int(m.group(2)), int(m.group(3)))

    # 26.9.6 のような2桁年。**1桁目だけでは日付と判断しない**（12.5 が数値のことがある）
    m = re.match(r"^(\d{2})[.\-/](\d{1,2})[.\-/](\d{1,2})$", s)
    if m:
        return _mk(2000 + int(m.group(1)), int(m.group(2)), int(m.group(3)))

    # 年が無い。文脈（ファイル名など）から補う。補えないなら諦める
    m = re.match(r"^(\d{1,2})[/\-月](\d{1,2})日?$", s)
    if m and default_year:
        return _mk(default_year, int(m.group(1)), int(m.group(2)))

    return None


def _mk(y, m, d):
    try:
        return datetime.date(y, m, d)
    except ValueError:
        return None


def parse_num(raw):
    u"""'１２．３' ' 12.3 ' '12.3℃' '約12' → 12.3。読めなければ None。"""
    if raw is None:
        return None
    s = z2h(str(raw)).strip()
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group(0)) if m else None


# --- 点検欄の判定 -----------------------------------------------------------

def classify(raw):
    u"""点検欄の1マスを6つに分ける。

    **空欄・判別不能・対象外・条件付き許容を、正常や異常に混ぜない。**

    判定の順番が本体。落とした誤りを2つ記録しておく。

      2026-09-12(1) **「異常なし」が「異常」で始まる**ため、接頭辞判定に先に当たって
        異常として数えられていた。正常71マスが全部 異常 に化ける。
        → **正常の語を先に見る。**接頭辞判定はその後にしか置けない。

      2026-09-12(2) レビュー指摘。**「要観察」を異常に数えていた。**
        現場では「一部劣化だが稼働継続」は異常ではない。異常件数が実態より膨らむ。
        また「計画停止中」の空欄を未記入に数えていた。**止まっていれば点検しない。**
        → 対象外 と 条件付き許容 を足して6分類に。
    """
    if raw is None or str(raw).strip() == u"":
        return BLANK, u""
    s = str(raw).strip()
    if z2h(s) in {z2h(x) for x in AMBIG_MARKS}:
        return AMBIGUOUS, s
    bare = z2h(s.rstrip(u"。"))
    if bare in {z2h(x.rstrip(u"。")) for x in OK_MARKS}:
        return RECORDED, s                      # ← 異常より先に見る
    if s.startswith(EXCLUDED_HEAD):
        return EXCLUDED, s
    if s.startswith(COND_HEAD):
        return CONDITIONAL, s                   # ← 異常より先に見る
    if s.startswith(ABNORMAL_HEAD):
        return ABNORMAL, s
    return RECORDED, s          # 自由記述。中身は本文にそのまま残す


def line_for(col, kind, raw):
    u"""1マスぶんの出力行。**状態ごとに、何が分からないのかまで書く。**"""
    if kind == BLANK:
        return u"- %s: **未記入**（点検されたかどうか、原票からは分からない）" % col
    if kind == AMBIGUOUS:
        return (u"- %s: **判別不能**（原文: %s）"
                u"— 「該当なし」か「未チェック」か決められない。**異常という意味ではない**"
                % (col, raw))
    if kind == EXCLUDED:
        return (u"- %s: **対象外** — %s"
                u"（設備が稼働していないため点検しない。未記入ではない）" % (col, raw))
    if kind == CONDITIONAL:
        return (u"- %s: **条件付き許容** — %s"
                u"（稼働は継続している。異常には数えない）" % (col, raw))
    if kind == ABNORMAL:
        return u"- %s: **異常** — %s" % (col, raw)
    return u"- %s: 記入あり — %s" % (col, raw)


# --- xlsx を読む ------------------------------------------------------------

def _col_index(ref):
    m = re.match(r"([A-Z]+)(\d+)", ref)
    col, n = m.group(1), 0
    for ch in col:
        n = n * 26 + (ord(ch) - 64)
    return n - 1, int(m.group(2))


def _have_openpyxl():
    try:
        import openpyxl
        return openpyxl.__version__
    except Exception:
        return None


def read_xlsx_builtin(path):
    u"""標準ライブラリだけで読む。**結合セルは左上の値で埋める。**

    埋めないと、日付が縦結合された帳票では3行のうち2行が「日付なし」になる。
    実務の帳票はほぼ必ずこれをやっている。
    """
    z = zipfile.ZipFile(path)

    shared = []
    if "xl/sharedStrings.xml" in z.namelist():
        root = ET.fromstring(z.read("xl/sharedStrings.xml"))
        for si in root.findall(NS + "si"):
            shared.append(u"".join(t.text or u"" for t in si.iter(NS + "t")))

    date_styles = set()
    if "xl/styles.xml" in z.namelist():
        st = ET.fromstring(z.read("xl/styles.xml"))
        custom = {int(n.get("numFmtId")): (n.get("formatCode") or u"")
                  for n in st.iter(NS + "numFmt")}
        xfs = st.find(NS + "cellXfs")
        if xfs is not None:
            for i, xf in enumerate(xfs.findall(NS + "xf")):
                fid = int(xf.get("numFmtId", "0"))
                if 14 <= fid <= 22 or re.search(r"[ymd]", custom.get(fid, u"")):
                    date_styles.add(i)

    sheet = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
    cells, maxc, maxr = {}, 0, 0
    for c in sheet.iter(NS + "c"):
        ref = c.get("r")
        if not ref:
            continue
        ci, ri = _col_index(ref)
        v = c.find(NS + "v")
        isel = c.find(NS + "is")
        if c.get("t") == "s" and v is not None:
            val = shared[int(v.text)]
        elif isel is not None:
            val = u"".join(t.text or u"" for t in isel.iter(NS + "t"))
        elif v is not None:
            val = v.text
            s = c.get("s")
            if s is not None and int(s) in date_styles:
                try:
                    val = serial_to_date(float(val)).isoformat()
                except Exception:
                    pass
        else:
            continue
        cells[(ri, ci)] = val
        maxr, maxc = max(maxr, ri), max(maxc, ci)

    merges = []
    for mc in sheet.iter(NS + "mergeCell"):
        a, b = mc.get("ref").split(":")
        (c0, r0), (c1, r1) = _col_index(a), _col_index(b)
        merges.append((r0, c0, r1, c1))
    for r0, c0, r1, c1 in merges:
        top = cells.get((r0, c0))
        if top in (None, u""):
            continue
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                cells.setdefault((r, c), top)

    return [[cells.get((r, c), u"") for c in range(maxc + 1)]
            for r in range(1, maxr + 1)], len(merges)


def _cellstr(v):
    if v is None:
        return u""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, float):
        return u"%g" % v
    return u"%s" % v


def read_xlsx_openpyxl(path):
    u"""openpyxl で読む。**あればこちらを優先する。**

    自前の読み取りは方言に弱い（1904年日付システム / Strict ネームスペース / 旧.xls）。
    **顧客の汚い Excel で確実に動くほうが、依存が無いことより価値がある。**
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = [[_cellstr(c.value) for c in row] for row in ws.iter_rows()]
    ranges = list(ws.merged_cells.ranges)
    for mr in ranges:
        r0, c0 = mr.min_row - 1, mr.min_col - 1
        if r0 >= len(rows) or c0 >= len(rows[r0]):
            continue
        top = rows[r0][c0]
        if top in (None, u""):
            continue
        for r in range(r0, min(mr.max_row, len(rows))):
            for c in range(c0, mr.max_col):
                if c < len(rows[r]) and rows[r][c] in (None, u""):
                    rows[r][c] = top
    return rows, len(ranges)


def read_xlsx(path, force_builtin=False):
    u"""返り値: (rows, 結合セル数, 使ったリーダー名)"""
    ver = None if force_builtin else _have_openpyxl()
    if ver:
        rows, n = read_xlsx_openpyxl(path)
        return rows, n, u"openpyxl %s" % ver
    rows, n = read_xlsx_builtin(path)
    return rows, n, u"内蔵（標準ライブラリのみ）"


def announce_reader(name):
    print(u"xlsx の読み取り: %s" % name)
    if name.startswith(u"内蔵"):
        print(u"  注意: 1904年日付システム / Strict ネームスペース / 旧.xls / 複数シート に未対応。"
              u"`pip install openpyxl` で差し替えられる")


# --- front-matter -----------------------------------------------------------
#
# ベクトル検索だけでは「2026年9月3日の」のような条件で絞れない。
# 意味の近さではなく**完全一致で絞るべきもの**がある（日付・設備・担当・部署）。
# → 各文書の先頭にこれを出し、rag_ja.py 側で metadata に入れてフィルタに使う。

FM_KEYS = (u"doc_type", u"date", u"month", u"equipment", u"person", u"dept", u"source", u"row")


def front_matter(**kv):
    L = [u"---"]
    for k in FM_KEYS:
        v = kv.get(k)
        if v in (None, u"", []):
            continue
        L.append(u"%s: %s" % (k, u",".join(v) if isinstance(v, list) else v))
    L.append(u"---")
    return u"\n".join(L) + u"\n\n"


# --- 点検記録 → Markdown ----------------------------------------------------

def norm_inspection(path, force_builtin=False):
    u"""点検記録を読み、**1点検＝1節**に分けて返す。

    返すもの: (header, sections, tally, per_col, unparsed, base, reader, month)
      sections: [(見出し, front-matter, 本文, ファイル名の芯), ...]  1点検（1行）が1つ
    """
    rows, n_merge, reader = read_xlsx(path, force_builtin)
    base = os.path.splitext(os.path.basename(path))[0]
    m = re.search(r"(\d{4})-(\d{2})", base)
    dy, dm = (int(m.group(1)), int(m.group(2))) if m else (None, None)
    month = u"%04d-%02d" % (dy, dm) if dy else u""

    head1, head2 = rows[0], rows[1]
    cols = [(z2h(a).strip() + u" " + z2h(b).strip()).strip()
            for a, b in zip(head1, list(head2) + [u""] * len(head1))]

    check_cols = [i for i, c in enumerate(cols)
                  if any(k in c for k in (u"掻き取り", u"異音", u"発泡", u"色"))]

    tally = dict((k, 0) for k in KINDS)
    per_col, sections, unparsed = {}, [], []

    header = (u"# 点検記録 " + base + u"（正規化済み）\n\n" +
              u"原票: `" + os.path.basename(path) + u"`（結合セル %d 箇所を展開／読み取り: %s）\n\n"
              % (n_merge, reader) +
              u"点検欄は6つに分けてある。"
              u"**未記入・判別不能・対象外・条件付き許容も、そのまま書き出している。**\n\n")

    for r, row in enumerate(rows[2:], start=3):
        if not any((u"%s" % x).strip() for x in row):
            continue
        d = parse_date(row[0], dy, dm)
        if d is None and (u"%s" % row[0]).strip():
            unparsed.append((r, row[0]))
        dstr = d.isoformat() if d else (u"日付が読めない（原文: %s）" % row[0])

        person = canon(u"担当", row[1])
        dept = canon(u"部署", row[7]) if len(row) > 7 else u""

        L = [u"- 日付: " + dstr,
             u"- 原票: %s の %d行目" % (os.path.basename(path), r),
             u"- 担当: %s（原文: %s）" % (person, row[1]),
             u"- 部署: %s（原文: %s）" % (dept, row[7] if len(row) > 7 else u"")]
        atsu = parse_num(row[2])
        L.append(u"- 汚泥厚: %s mm（原文: %s）"
                 % (u"%.1f" % atsu if atsu is not None else u"読めない",
                    (u"%s" % row[2]).strip()))

        equip = []
        for i in check_cols:
            kind, raw = classify(row[i] if i < len(row) else u"")
            tally[kind] += 1
            per_col.setdefault(cols[i], dict((k, 0) for k in KINDS))
            per_col[cols[i]][kind] += 1
            L.append(line_for(cols[i], kind, raw))
            equip += find_equipment(cols[i])
        if len(row) > 8 and (u"%s" % row[8]).strip():
            L.append(u"- 備考: " + (u"%s" % row[8]).strip())

        key = d.isoformat() if d else u"日付不明"
        fm = front_matter(doc_type=u"inspection", date=(d.isoformat() if d else u""),
                          month=month, equipment=sorted(set(equip)),
                          person=person, dept=dept,
                          source=os.path.basename(path), row=u"%d" % r)
        sections.append((u"%s  %s の点検（%d行目）" % (key, base, r),
                         fm, u"\n".join(L), u"%s_r%02d" % (key, r)))

    return header, sections, tally, per_col, unparsed, base, reader, month


def summary_doc(base, tally, per_col, unparsed, month):
    u"""**集計は文書に書く。**RAG は検索であって集計ではない。"""
    total = sum(tally.values())
    checked = sum(tally[k] for k in COUNTED_AS_CHECKED)
    L = [front_matter(doc_type=u"summary", month=month, source=base),
         u"# 点検記録 " + base + u" の記入状況（集計）\n",
         u"点検欄 %d マスの内訳。\n" % total]
    for k in KINDS:
        L.append(u"- %s: %d マス（%.1f%%）" % (k, tally[k], 100.0 * tally[k] / total))
    L.append(u"")
    L.append(u"**記入率は %.1f%%**"
             u"（記入あり・異常・条件付き許容・対象外を、点検された記録として数えた）。\n"
             % (100.0 * checked / total))
    L.append(u"**未記入 %d マスと判別不能 %d マスは、異常が無かったことを意味しない。**"
             u"点検された記録が無いだけ。合わせて %d マス。\n"
             % (tally[BLANK], tally[AMBIGUOUS], tally[BLANK] + tally[AMBIGUOUS]))
    L.append(u"**対象外 %d マスは未記入ではない。**設備が稼働していないため点検しない。\n"
             % tally[EXCLUDED])
    L.append(u"**条件付き許容 %d マスは異常ではない。**稼働は継続している。\n" % tally[CONDITIONAL])
    L.append(u"## 点検項目ごと\n")
    for col, t in per_col.items():
        n = sum(t.values())
        c = sum(t[k] for k in COUNTED_AS_CHECKED)
        L.append(u"- %s: 記入あり %d / 異常 %d / 条件付き許容 %d / 対象外 %d / 未記入 %d / 判別不能 %d"
                 u"（記入率 %.0f%%）"
                 % (col, t[RECORDED], t[ABNORMAL], t[CONDITIONAL], t[EXCLUDED],
                    t[BLANK], t[AMBIGUOUS], 100.0 * c / n))
    if unparsed:
        L.append(u"\n## 日付が読めなかった行\n")
        for r, raw in unparsed:
            L.append(u"- %d行目: `%s`" % (r, raw))
    else:
        L.append(u"\n## 日付が読めなかった行\n\n- なし")
    return u"\n".join(L)


def norm_nippou(path, text):
    base = os.path.splitext(os.path.basename(path))[0]
    d = None
    first = text.strip().split(u"\n")[0]
    m = re.match(r"^(\S+)\s*日報", first)
    if m:
        d = parse_date(m.group(1), 2026, 9)
    body, person, dept = [], u"", u""
    for line in text.strip().split(u"\n")[1:]:
        s = line.strip()
        if s.startswith(u"部署:"):
            v = s.split(u":", 1)[1].strip()
            dept = canon(u"部署", v)
            body.append(u"- 部署: %s（原文: %s）" % (dept, v))
        elif s.startswith(u"記入:"):
            v = s.split(u":", 1)[1].strip()
            person = canon(u"担当", v)
            body.append(u"- 担当: %s（原文: %s）" % (person, v))
        else:
            body.append(canon_text(u"設備", s))
    text2 = u"\n".join(body)
    fm = front_matter(doc_type=u"nippou", date=(d.isoformat() if d else u""),
                      month=(d.isoformat()[:7] if d else u""),
                      equipment=find_equipment(text2), person=person, dept=dept,
                      source=os.path.basename(path))
    head = u"# 日報 " + (d.isoformat() if d else base) + u"\n"
    if d is None:
        head += u"\n（日付が読めない。原文: " + first + u"）\n"
    return fm + head + u"\n" + text2 + u"\n"


# --- 対照（正規化しない読み方）----------------------------------------------

def naive_inspection(path):
    u"""**素朴に読んだらどうなるか。**効果を測るための対照。

    素朴な実装がやること（どれも悪意はない。普通に書くとこうなる）:
      - 結合セルを展開しない  → 日付が入っているのは先頭行だけ
      - 日付をそのまま文字列で出す
      - 表記ゆれをそのままにする
      - **空欄のセルは行に出さない**（値が無いのだから書きようがない）

    3つ目が効く。**空欄が消えるので、残るのは「異常なし」だけになる。**
    """
    z = zipfile.ZipFile(path)
    shared = []
    if "xl/sharedStrings.xml" in z.namelist():
        root = ET.fromstring(z.read("xl/sharedStrings.xml"))
        for si in root.findall(NS + "si"):
            shared.append(u"".join(t.text or u"" for t in si.iter(NS + "t")))
    sheet = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))

    rows = {}
    for c in sheet.iter(NS + "c"):
        ref = c.get("r")
        if not ref:
            continue
        ci, ri = _col_index(ref)
        v = c.find(NS + "v")
        if v is None:
            continue
        val = shared[int(v.text)] if c.get("t") == "s" else v.text
        rows.setdefault(ri, {})[ci] = val

    base = os.path.splitext(os.path.basename(path))[0]
    out = [u"# " + base, u""]
    n_cells = 0
    for ri in sorted(rows):
        if ri <= 2:
            continue
        cells = rows[ri]
        n_cells += sum(1 for ci in (3, 4, 5, 6) if ci in cells)
        out.append(u"| " + u" | ".join(cells.get(ci, u"") for ci in sorted(cells)) + u" |")
    return u"\n".join(out), n_cells


def run_naive(src, dst):
    os.makedirs(dst, exist_ok=True)
    n_cells = 0
    for dirpath, dirnames, filenames in os.walk(src):
        for fn in sorted(filenames):
            p = os.path.join(dirpath, fn)
            if fn.lower().endswith(".xlsx") and not fn.startswith("~$"):
                md, n = naive_inspection(p)
                n_cells += n
                io.open(os.path.join(dst, os.path.splitext(fn)[0] + u".md"),
                        "w", encoding="utf-8").write(md)
            elif fn.lower().endswith((".txt", ".md")):
                io.open(os.path.join(dst, os.path.splitext(fn)[0] + u".md"),
                        "w", encoding="utf-8").write(
                    io.open(p, encoding="utf-8").read())
    print(u"対照（正規化なし）: " + dst)
    print(u"  点検欄として出てきたマス: %d" % n_cells)
    print(u"  **差は、空欄が消えたぶん。**")
    return 0


# --- 入口 -------------------------------------------------------------------

def run(src, dst, per_day=True, force_builtin=False):
    os.makedirs(dst, exist_ok=True)
    n_x = n_t = n_sec = 0
    for dirpath, dirnames, filenames in os.walk(src):
        for fn in sorted(filenames):
            p = os.path.join(dirpath, fn)
            if fn.lower().endswith(".xlsx") and not fn.startswith("~$"):
                (header, sections, tally, per_col,
                 unp, base, reader, month) = norm_inspection(p, force_builtin)
                announce_reader(reader)
                if per_day:
                    # **1点検＝1文書。**引用が「どの日のどの行か」まで絞れる
                    sub = os.path.join(dst, base)
                    os.makedirs(sub, exist_ok=True)
                    for title, fm, body, stem in sections:
                        io.open(os.path.join(sub, stem + u".md"), "w",
                                encoding="utf-8").write(
                            fm + u"# " + title + u"\n\n" + body + u"\n")
                else:
                    io.open(os.path.join(dst, base + u".md"), "w",
                            encoding="utf-8").write(
                        header + u"\n".join(u"## " + t + u"\n\n" + b + u"\n"
                                            for t, _f, b, _s in sections))
                io.open(os.path.join(dst, base + u"_集計.md"), "w",
                        encoding="utf-8").write(
                    summary_doc(base, tally, per_col, unp, month))
                n_x += 1
                n_sec += len(sections)
                print(u"点検記録: " + fn)
                print(u"  " + u" / ".join(u"%s %d" % (k, tally[k]) for k in KINDS))
                print(u"  日付が読めなかった行: %d" % len(unp))
            elif fn.lower().endswith((".txt", ".md")):
                text = io.open(p, encoding="utf-8").read()
                io.open(os.path.join(dst, os.path.splitext(fn)[0] + u".md"),
                        "w", encoding="utf-8").write(norm_nippou(p, text))
                n_t += 1
    print(u"出力: %s（点検記録 %d → %s / 日報 %d）"
          % (dst, n_x,
             (u"%d文書（1点検=1文書）" % n_sec) if per_day else u"1文書", n_t))
    print(u"次: python rag_ja.py index " + dst)
    return 0


def _cross_check():
    u"""**内蔵の読み取りと openpyxl を突き合わせる。**

    自前で xlsx を解くのは方言に弱い。弱いことは認めたうえで、
    **「手元のファイルでは openpyxl と同じ答えを返す」**ことは示せる。
    openpyxl が無い環境では黙って飛ばす（それも表示する）。
    """
    ver = _have_openpyxl()
    sample = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                          u"samples", u"raw", u"点検記録_2026-09.xlsx")
    if not ver:
        print(u"  突き合わせ: openpyxl が無いので飛ばした")
        return
    if not os.path.exists(sample):
        print(u"  突き合わせ: サンプルが無いので飛ばした（python samples/make_samples.py）")
        return
    a, na = read_xlsx_builtin(sample)
    b, nb = read_xlsx_openpyxl(sample)
    assert na == nb, u"結合セルの数が違う: 内蔵 %d / openpyxl %d" % (na, nb)
    diff = 0
    for r in range(max(len(a), len(b))):
        ra = a[r] if r < len(a) else []
        rb = b[r] if r < len(b) else []
        for c in range(max(len(ra), len(rb))):
            va = z2h((u"%s" % ra[c]).strip()) if c < len(ra) else u""
            vb = z2h((u"%s" % rb[c]).strip()) if c < len(rb) else u""
            if va != vb:
                diff += 1
                if diff <= 3:
                    print(u"  ちがい r%dc%d: 内蔵=%r / openpyxl=%r" % (r + 1, c + 1, va, vb))
    assert diff == 0, u"内蔵と openpyxl で %d マス違う" % diff
    print(u"  突き合わせ: 内蔵 == openpyxl %s（結合%d箇所・全マス一致）" % (ver, na))


def demo():
    u"""自己確認。**壊れたら落ちる。**"""
    D = datetime.date
    assert parse_date(u"2026/9/1") == D(2026, 9, 1)
    assert parse_date(u"2026-09-02") == D(2026, 9, 2)
    assert parse_date(u"令和8年9月3日") == D(2026, 9, 3)
    assert parse_date(u"R8.9.5") == D(2026, 9, 5)
    assert parse_date(u"26.9.6") == D(2026, 9, 6)
    assert parse_date(u"9/4", 2026, 9) == D(2026, 9, 4)
    assert parse_date(u"9/4") is None, u"年が補えないなら諦める（でっち上げない）"
    assert parse_date(u"46266") == D(2026, 9, 1), u"Excel シリアル値"
    assert parse_date(u"") is None and parse_date(None) is None
    assert parse_date(u"12.5") is None, u"2桁.1桁 は数値かもしれない。日付と決めつけない"
    assert parse_date(u"R8.9") is None, u"3要素そろわなければ日付にしない"
    assert parse_date(u"V2.1.0") is None and parse_date(u"LOT-25.3.1") is None
    assert parse_date(u"99.99.99") is None, u"日付として成立しない"
    assert parse_date(u"12.3.4") == D(2012, 3, 4), u"**既知の誤検知。**日付列以外から呼ばない"

    assert parse_num(u"１２．３") == 12.3
    assert parse_num(u"  12.3 ") == 12.3
    assert parse_num(u"12.3℃") == 12.3
    assert parse_num(u"約12") == 12.0
    assert parse_num(u"") is None

    assert canon(u"部署", u"Ｍ１") == u"第1製造部"
    assert canon(u"部署", u"第一製造部") == u"第1製造部"
    assert canon(u"担当", u"ヤマダ") == u"山田太郎"
    assert canon(u"担当", u"鈴木") == u"鈴木", u"表に無ければそのまま返す"

    # 本文の置換。**既に正しい形に二度当てない**（2026-09-12 に実データで出た）
    assert canon_text(u"設備", u"加圧浮上槽の点検") == u"加圧浮上槽の点検"
    assert canon_text(u"設備", u"浮上槽の点検") == u"加圧浮上槽の点検"
    assert canon_text(u"設備", u"加圧浮上の点検") == u"加圧浮上槽の点検"
    assert canon_text(u"設備", u"活性汚泥槽と汚泥槽") == u"活性汚泥槽と活性汚泥槽"
    assert find_equipment(u"浮上槽と汚泥槽を見た") == [u"加圧浮上槽", u"活性汚泥槽"]

    # 点検欄の6分類
    assert classify(u"異常なし")[0] == RECORDED, u"「異常」で始まるが正常（2026-09-12 に取りこぼした）"
    assert classify(u"異常無し")[0] == RECORDED
    assert classify(u"異常なし。")[0] == RECORDED
    assert classify(u"✓")[0] == RECORDED
    assert classify(u"")[0] == BLANK
    assert classify(None)[0] == BLANK
    assert classify(u"   ")[0] == BLANK
    assert classify(u"-")[0] == AMBIGUOUS
    assert classify(u"ー")[0] == AMBIGUOUS
    assert classify(u"／")[0] == AMBIGUOUS
    assert classify(u"N/A")[0] == AMBIGUOUS
    assert classify(u"異常あり：掻き取り機に異音")[0] == ABNORMAL
    assert classify(u"NG")[0] == ABNORMAL
    assert classify(u"不良")[0] == ABNORMAL
    # レビュー指摘で足した2つ（2026-09-12）
    assert classify(u"要観察（発泡多め）")[0] == CONDITIONAL, u"**異常ではない。**稼働は継続している"
    assert classify(u"経過観察")[0] == CONDITIONAL
    assert classify(u"軽微な劣化")[0] == CONDITIONAL
    assert classify(u"要確認")[0] == CONDITIONAL
    assert classify(u"計画停止中")[0] == EXCLUDED, u"**未記入ではない。**止まっていれば点検しない"
    assert classify(u"休転")[0] == EXCLUDED
    assert classify(u"対象外")[0] == EXCLUDED
    assert classify(u"非稼働")[0] == EXCLUDED

    # 出力行に、その状態が何を意味しないのかまで書く
    assert u"異常という意味ではない" in line_for(u"色", AMBIGUOUS, u"-")
    assert u"未記入ではない" in line_for(u"色", EXCLUDED, u"計画停止")
    assert u"異常には数えない" in line_for(u"色", CONDITIONAL, u"要観察")

    fm = front_matter(doc_type=u"inspection", date=u"2026-09-03",
                      equipment=[u"加圧浮上槽"], person=u"山田太郎", row=u"6")
    assert fm.startswith(u"---") and u"date: 2026-09-03" in fm and u"row: 6" in fm
    assert u"month:" not in fm, u"空の項目は出さない"

    _cross_check()
    print(u"demo: ok（日付 / 数値 / 表記ゆれ / 点検欄6分類 / front-matter / 読み取りの突き合わせ）")
    return 0


def main(argv=None):
    ap = argparse.ArgumentParser(description=u"現場の帳票を RAG が読める Markdown にする")
    ap.add_argument("src", nargs="?", help=u"入力ディレクトリ（.xlsx / .txt）")
    ap.add_argument("-o", "--out", default="_norm", help=u"出力ディレクトリ")
    ap.add_argument("--demo", action="store_true", help=u"自己確認")
    # 既定を「1点検=1文書」にしたのは測定の結果（MEASURE_GENBA.md §5）。
    # 1ファイルにまとめると、特定日の点検が上位5件に入らなかった。
    ap.add_argument("--single-file", dest="single", action="store_true",
                    help=u"点検記録を1ファイルにまとめる（既定は1点検=1文書）")
    ap.add_argument("--builtin-reader", dest="builtin", action="store_true",
                    help=u"openpyxl があっても内蔵の読み取りを使う（比較用）")
    ap.add_argument("--naive", action="store_true",
                    help=u"正規化しない対照を出す（効果を測るため）")
    a = ap.parse_args(argv)
    if a.demo:
        return demo()
    if not a.src:
        ap.print_help()
        return 2
    return run_naive(a.src, a.out) if a.naive else run(a.src, a.out, not a.single, a.builtin)


if __name__ == "__main__":
    sys.exit(main())
